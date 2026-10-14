"""File bytes → structured text blocks (headings, paragraphs, tables, page numbers)."""
from __future__ import annotations

import csv
import io
import re
import zipfile
from html.parser import HTMLParser
from typing import Literal

from app.rag.models import Block, LoadedDoc

Kind = Literal["pdf", "docx", "xlsx", "html", "markdown", "csv", "text"]
ROWS_PER_TABLE = 50
TEXT_EXTENSIONS = {
    ".txt", ".json", ".log", ".yaml", ".yml", ".xml", ".py", ".ts", ".tsx", ".js", ".jsx", ".sql", ".java", ".go",
    ".rs", ".c", ".cpp", ".h", ".cs", ".rb", ".php", ".sh", ".ps1", ".toml", ".ini", ".cfg", ".rst", ".tex",
}


MAX_UNCOMPRESSED = 1024 * 1024 * 1024  # 1 GB: a 25 MB upload can't honestly reach this; a bomb declares many GB
MAX_RATIO = 200  # real Office files expand ~5-30x; 200x+ is a compression bomb


class UnsupportedFile(ValueError):
    """The file can't be added; the message is shown to the user."""


def _guard_zip(data: bytes) -> None:
    """Refuse a docx/xlsx (zip) that declares more uncompressed bytes than we would ever read, before touching a
    member: a few-KB upload can expand to gigabytes and OOM the shared ingest worker.
    ponytail: trusts the zip directory's sizes (honest in real-world bombs); the container memory limit is the backstop."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile:
        return  # not a real zip; detect_kind reports the unsupported type
    if total > MAX_UNCOMPRESSED or total > max(len(data), 1) * MAX_RATIO:
        raise UnsupportedFile(f"This file is too large to open (it expands to about {total // (1024 * 1024)} MB).")


def detect_kind(name: str, data: bytes) -> Kind:
    if not data.strip():
        raise UnsupportedFile(f"{name} is empty")
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        _guard_zip(data)
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                types = archive.read("[Content_Types].xml").decode("utf-8", "ignore")
        except (zipfile.BadZipFile, KeyError):
            types = ""
        if "wordprocessingml" in types:
            return "docx"
        if "spreadsheetml" in types:
            return "xlsx"
        raise UnsupportedFile(f"{name}: this file type isn't supported (archives and other Office files can't be added)")
    if b"\x00" in data[:4096] or data.startswith((b"\x89PNG", b"MZ", b"\xff\xd8\xff", b"GIF8")):
        raise UnsupportedFile(f"{name}: this file type isn't supported (only text, PDF, Word and Excel files)")
    if ext in (".html", ".htm"):
        return "html"
    if ext in (".md", ".markdown"):
        return "markdown"
    if ext == ".csv":
        return "csv"
    return "text"


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _paragraphs(text: str, page: int | None = None) -> list[Block]:
    return [Block(p.strip(), "paragraph", page=page) for p in re.split(r"\n\s*\n", text) if p.strip()]


def _table_blocks(rows: list[list[str]]) -> list[Block]:
    """Row groups of ROWS_PER_TABLE with the header row repeated, cells joined with " | "."""
    rows = [r for r in rows if any(cell.strip() for cell in r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:] or [[]]
    blocks = []
    for start in range(0, len(body), ROWS_PER_TABLE):
        group = [header] + body[start:start + ROWS_PER_TABLE]
        blocks.append(Block("\n".join(" | ".join(c.strip() for c in row) for row in group if row), "table"))
    return blocks


def _markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    buffer: list[str] = []

    def flush():
        if buffer:
            blocks.extend(_paragraphs("\n".join(buffer)))
            buffer.clear()

    for line in text.splitlines():
        heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
        if heading:
            flush()
            blocks.append(Block(heading.group(2), "heading", level=len(heading.group(1))))
        else:
            buffer.append(line)
    flush()
    return blocks


class _HTMLBlocks(HTMLParser):
    SKIP = {"script", "style", "noscript", "head"}
    BREAKS = {"p", "div", "li", "br", "tr", "section", "article", "table", "ul", "ol", "pre", "blockquote"}

    def __init__(self) -> None:
        super().__init__()
        self.blocks: list[Block] = []
        self.text: list[str] = []
        self.skipping = 0
        self.heading = 0

    def _flush(self) -> None:
        text = " ".join("".join(self.text).split())
        if text:
            self.blocks.append(Block(text, "heading", level=self.heading) if self.heading else Block(text, "paragraph"))
        self.text = []

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skipping += 1
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = int(tag[1])
        elif tag in self.BREAKS:
            self._flush()

    def handle_endtag(self, tag):
        if tag in self.SKIP:
            self.skipping = max(0, self.skipping - 1)
        elif re.fullmatch(r"h[1-6]", tag):
            self._flush()
            self.heading = 0
        elif tag in self.BREAKS:
            self._flush()

    def handle_data(self, data):
        if not self.skipping:
            self.text.append(data)


def _pdf(name: str, data: bytes) -> LoadedDoc:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise UnsupportedFile(f"{name}: the PDF could not be read ({exc})") from exc
    blocks = [block for number, text in enumerate(pages, start=1) for block in _paragraphs(text, page=number)]
    if not blocks:
        raise UnsupportedFile(f"{name} has no selectable text (it may be a scanned image)")
    return LoadedDoc(blocks, len(pages))


def _docx(data: bytes) -> list[Block]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    blocks: list[Block] = []
    # Walk paragraphs and tables in document order.
    tables = {table._tbl: table for table in doc.tables}
    paragraphs = {p._p: p for p in doc.paragraphs}
    for element in doc.element.body.iterchildren():
        if element in paragraphs:
            p = paragraphs[element]
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            level = re.match(r"Heading (\d)", style)
            blocks.append(Block(text, "heading", level=int(level.group(1))) if level else
                          Block(text, "heading", level=1) if style == "Title" else Block(text, "paragraph"))
        elif element in tables:
            rows = [[cell.text for cell in row.cells] for row in tables[element].rows]
            blocks.extend(_table_blocks(rows))
    return blocks


def _xlsx(data: bytes) -> list[Block]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    blocks: list[Block] = []
    for sheet in book.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        tables = _table_blocks(rows)
        if tables:
            blocks.append(Block(sheet.title, "heading", level=1))
            blocks.extend(tables)
    return blocks


def load(name: str, data: bytes) -> LoadedDoc:
    kind = detect_kind(name, data)
    if kind == "pdf":
        return _pdf(name, data)
    if kind == "docx":
        blocks = _docx(data)
    elif kind == "xlsx":
        blocks = _xlsx(data)
    else:
        text = _decode(data)
        if kind == "html":
            parser = _HTMLBlocks()
            parser.feed(text)
            parser._flush()
            blocks = parser.blocks
        elif kind == "markdown":
            blocks = _markdown(text)
        elif kind == "csv":
            sample = text[:4096]
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
            blocks = _table_blocks(list(csv.reader(io.StringIO(text), delimiter=delimiter)))
        else:
            blocks = _paragraphs(text)
    if not blocks:
        raise UnsupportedFile(f"{name} has no text")
    return LoadedDoc(blocks, None)
