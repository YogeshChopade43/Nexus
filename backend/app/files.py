"""Files the agent makes: CSV, Excel, Word, Markdown, text, JSON or HTML, to download from the chat, receive on
Telegram, or attach to an email. They live next to the database, under files/<id>/<name>."""
from __future__ import annotations

import csv
import io
import json
import re
import secrets
from pathlib import Path
from typing import Any

from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import FileResponse

from app.core.database import get_db_path

FORMATS = {".csv", ".xlsx", ".docx", ".md", ".txt", ".json", ".html"}
MAX_CONTENT = 2_000_000  # characters the model may put into one file
MAX_UPLOAD = 20_000_000  # bytes of one file attached in the chat
PREVIEW_ROWS = 200  # rows of each sheet a preview shows
PREVIEW_CHARS = 200_000  # characters of a text preview

router = APIRouter(prefix="/api/files", tags=["files"])


def _root() -> Path:
    return get_db_path().parent / "files"


def _safe_name(filename: str) -> str:
    name = re.sub(r"[^\w.\- ]+", "_", Path(filename.replace("\\", "/")).name).strip(" ._") or "file.txt"
    if Path(name).suffix.lower() not in FORMATS:
        raise ValueError(f"Nexus can't make that kind of file: use one of {', '.join(sorted(FORMATS))}")
    return name[:120]


def _markdown_table(content: str) -> list[list[str]] | None:
    """Rows of a Markdown table (``| a | b |``, models' favourite way to write one), without the ``|---|`` line."""
    lines = [line.strip() for line in content.strip().splitlines() if line.strip()]
    if not lines or not all(line.startswith("|") for line in lines):
        return None
    return [[cell.strip() for cell in line.strip("|").split("|")] for line in lines
            if not re.fullmatch(r"[|:\-\s]+", line)]


def _rows(content: str) -> list[list[Any]]:
    """A table from JSON (a list of objects or of lists), a Markdown table, or CSV text."""
    try:
        data = json.loads(content)
    except ValueError:
        return _markdown_table(content) or list(csv.reader(io.StringIO(content)))
    if isinstance(data, list) and data and all(isinstance(row, dict) for row in data):
        header = list(dict.fromkeys(key for row in data for key in row))
        return [header, *[[row.get(key, "") for key in header] for row in data]]
    if isinstance(data, list) and all(isinstance(row, list) for row in data):
        return data
    raise ValueError("For a table, give a JSON list of objects or of rows, or CSV text")


def _cell(value: Any) -> Any:
    return value if isinstance(value, (int, float, str)) or value is None else json.dumps(value)


def _number(value: Any) -> Any:
    """"12" and "60.5" as numbers, so Excel can sum them; "007" and dates stay text."""
    if isinstance(value, str) and re.fullmatch(r"-?(0|[1-9]\d*)(\.\d+)?", value.strip()):
        return float(value) if "." in value else int(value)
    return value


def _xlsx(rows: list[list[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    for row in rows:
        sheet.append([_number(_cell(v)) for v in row])
    if rows:  # a header that looks like one, stays in view, and filters; columns wide enough to read
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="E8E6F0")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(sheet.iter_cols(), start=1):
            longest = max(len(str(c.value)) for c in column if c.value is not None) if any(c.value is not None for c in column) else 0
            sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 8), 60)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _inline(paragraph: Any, text: str) -> None:
    """**bold** and *italic* runs in a Word paragraph."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part:
            paragraph.add_run(part)


def _docx(markdown: str) -> bytes:
    from docx import Document

    doc = Document()
    table: list[str] = []
    for line in [*markdown.splitlines(), ""]:  # the extra line ends a table at the very end
        stripped = line.strip()
        if stripped.startswith("|"):
            table.append(stripped)
            continue
        if table:
            rows, table = _markdown_table("\n".join(table)) or [], []
            if rows:
                grid = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows), style="Table Grid")
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        grid.cell(r, c).text = value
                        if r == 0:
                            for run in grid.cell(r, c).paragraphs[0].runs:
                                run.bold = True
        heading = re.match(r"^(#{1,4})\s+(.*)", stripped)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
        elif re.match(r"^[-*]\s+", stripped):
            _inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+[.)]\s+", stripped):
            _inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+[.)]\s+", "", stripped))
        elif stripped:
            _inline(doc.add_paragraph(), stripped)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def render(name: str, content: str) -> bytes:
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        if content.lstrip().startswith(("[", "{", "|")):
            out = io.StringIO()
            csv.writer(out, lineterminator="\n").writerows([[_cell(v) for v in row] for row in _rows(content)])
            content = out.getvalue().rstrip("\n")
        return "﻿".encode() + content.encode("utf-8")  # the BOM makes Excel read UTF-8 correctly
    if suffix == ".xlsx":
        return _xlsx(_rows(content))
    if suffix == ".docx":
        return _docx(content)
    return content.encode("utf-8")


def _store(name: str, data: bytes) -> dict[str, Any]:
    file_id = secrets.token_hex(8)
    folder = _root() / file_id
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_bytes(data)
    return {"id": file_id, "name": name, "size": len(data), "url": f"/api/files/{file_id}/{name}"}


def create_file(filename: str, content: str) -> dict[str, Any]:
    name = _safe_name(filename)
    if len(content) > MAX_CONTENT:
        raise ValueError("That's too much content for one file")
    return _store(name, render(name, content))


def _path(file_id: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{16}", file_id or ""):
        raise ValueError(f"No file {file_id}")
    folder = _root() / file_id
    found = [p for p in folder.iterdir() if p.is_file()] if folder.is_dir() else []
    if not found:
        raise ValueError(f"No file {file_id}")
    return found[0]


def attachments(ids: str) -> list[tuple[str, bytes]]:
    """(name, bytes) for each file id in a comma-separated list (from create_file)."""
    out = []
    for file_id in (i.strip() for i in (ids or "").split(",")):
        if file_id:
            path = _path(file_id)
            out.append((path.name, path.read_bytes()))
    return out


def purge_older_than(days: int) -> int:
    """Delete created files older than ``days`` (0 keeps everything). Returns how many went."""
    import shutil
    import time

    if days <= 0 or not _root().is_dir():
        return 0
    cutoff, removed = time.time() - days * 86400, 0
    for folder in _root().iterdir():
        if folder.is_dir() and folder.stat().st_mtime < cutoff:
            shutil.rmtree(folder, ignore_errors=True)
            removed += 1
    return removed


def attach_to(msg: Any, ids: str) -> None:
    """Add the created files ``ids`` (comma separated) to an EmailMessage whose body is already set."""
    import mimetypes

    for name, data in attachments(ids):
        kind = (mimetypes.guess_type(name)[0] or "application/octet-stream").split("/", 1)
        msg.add_attachment(data, maintype=kind[0], subtype=kind[1], filename=name)


# ---- previews -----------------------------------------------------------------------------------------

def _plain(value: Any) -> Any:
    return "" if value is None else value if isinstance(value, (int, float, str)) else str(value)


def _sheets(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        named: list[tuple[str, Any]] = [("Sheet", csv.reader(io.StringIO(path.read_text(encoding="utf-8-sig", errors="replace"))))]
    else:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        named = [(sheet.title, sheet.iter_rows(values_only=True)) for sheet in book.worksheets[:5]]
    sheets = []
    for title, rows in named:
        shown, truncated = [], False
        for row in rows:
            if len(shown) == PREVIEW_ROWS:
                truncated = True
                break
            shown.append([_plain(v) for v in list(row)[:30]])
        sheets.append({"name": title, "rows": shown, "truncated": truncated})
    return sheets


def _docx_markdown(path: Path) -> str:
    """A Word document as Markdown: headings, lists, paragraphs and tables, in order."""
    from docx import Document
    from docx.table import Table

    blocks = []
    for block in Document(str(path)).iter_inner_content():
        if isinstance(block, Table):
            rows = [[c.text.strip().replace("|", r"\|").replace("\n", " ") for c in row.cells] for row in block.rows]
            if rows:
                blocks.append("\n".join([f"| {' | '.join(rows[0])} |", "|" + "---|" * len(rows[0]),
                                         *[f"| {' | '.join(r)} |" for r in rows[1:]]]))
            continue
        style, text = block.style.name if block.style else "", block.text.strip()
        heading = re.match(r"Heading (\d)", style)
        if text:
            blocks.append(f"{'#' * int(heading.group(1))} {text}" if heading else
                          f"- {text}" if "Bullet" in style else f"1. {text}" if "Number" in style else text)
    return "\n\n".join(blocks)


def _text(text: str, kind: str = "text") -> dict[str, Any]:
    return {"kind": kind, "text": text[:PREVIEW_CHARS], "truncated": len(text) > PREVIEW_CHARS}


def preview(path: Path) -> dict[str, Any]:
    """What the chat shows for a file: a table, Markdown, text, a web page (shown sandboxed), a PDF, or nothing."""
    suffix = path.suffix.lower()
    if suffix in (".csv", ".xlsx"):
        return {"kind": "table", "sheets": _sheets(path)}
    if suffix == ".docx":
        return _text(_docx_markdown(path), "markdown")
    if suffix == ".pdf":
        return {"kind": "pdf"}
    try:
        text = path.read_bytes().decode("utf-8-sig")
    except UnicodeDecodeError:
        return {"kind": "none"}
    if "\x00" in text:
        return {"kind": "none"}
    if suffix == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except ValueError:
            pass
    return _text(text, {".md": "markdown", ".html": "html"}.get(suffix, "text"))


def _file_or_404(file_id: str, name: str) -> Path:
    try:
        path = _path(file_id)
    except ValueError:
        raise HTTPException(status_code=404, detail="File not found")
    if path.name != name:
        raise HTTPException(status_code=404, detail="File not found")
    return path


@router.post("")
async def upload(request: Request) -> list[dict[str, Any]]:
    """Files attached in the chat, kept so the chat can preview them (multipart field ``file``)."""
    form = await request.form()
    uploads = [f for f in form.getlist("file") if hasattr(f, "read")]
    for item in uploads:  # refuse before reading: a huge file must not be loaded into memory
        if item.size is not None and item.size > MAX_UPLOAD:
            raise HTTPException(status_code=422, detail=f"{item.filename} is larger than 20 MB")
    stored = []
    for item in uploads:
        name = re.sub(r"[^\w.\- ]+", "_", Path((item.filename or "").replace("\\", "/")).name).strip(" ._")[:120] or "file"
        stored.append(_store(name, await item.read()))
    return stored


@router.get("/{file_id}/{name}/preview")
def file_preview(file_id: str, name: str) -> dict[str, Any]:
    path = _file_or_404(file_id, name)
    try:
        return preview(path)
    except Exception:  # a damaged spreadsheet or document: the download still works
        return {"kind": "none"}


@router.get("/{file_id}/{name}")
def download(file_id: str, name: str, inline: bool = False) -> FileResponse:
    path = _file_or_404(file_id, name)
    # Only PDFs open in the page (the browser's own viewer); anything else, HTML above all, is a download.
    shown = inline and path.suffix.lower() == ".pdf"
    return FileResponse(path, filename=path.name, content_disposition_type="inline" if shown else "attachment")
