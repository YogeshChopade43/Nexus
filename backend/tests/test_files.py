"""Files out: the agent creates CSV / Excel / Word / text files to download, send on Telegram or attach to email."""
from __future__ import annotations

import asyncio
import io
import json

import pytest
from fastapi.testclient import TestClient

from app import files
from app.main import app
from app.mcp import server_states
from app.tools import get_tool_registry

ROWS = json.dumps([{"title": "Login bug", "state": "open"}, {"title": "Typo", "state": "closed"}])


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def create(**parameters):
    return asyncio.run(get_tool_registry().get_tool("files__create_file").execute(parameters))["data"]["file"]


def download(client, made) -> bytes:
    response = client.get(made["url"])
    assert response.status_code == 200, response.text
    return response.content


def test_csv_from_rows_or_text(client):
    made = create(filename="issues.csv", content=ROWS)
    assert made["name"] == "issues.csv" and made["url"] == f"/api/files/{made['id']}/issues.csv"
    assert download(client, made).decode("utf-8-sig").splitlines() == ["title,state", "Login bug,open", "Typo,closed"]
    assert download(client, create(filename="raw.csv", content="a,b\n1,2")).decode("utf-8-sig") == "a,b\n1,2"


def test_excel_and_word(client):
    from docx import Document
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(download(client, create(filename="issues.xlsx", content=ROWS)))).active
    assert [c.value for c in sheet[1]] == ["title", "state"] and sheet["A3"].value == "Typo"

    doc = Document(io.BytesIO(download(client, create(filename="report.docx",
                                                       content="# Weekly report\n\n- **Shipped** login fix\n- Typo\n\nAll good."))))
    paragraphs = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert paragraphs[0] == ("Heading 1", "Weekly report")
    assert ("List Bullet", "Shipped login fix") in paragraphs and paragraphs[-1][1] == "All good."
    assert any(run.bold for p in doc.paragraphs for run in p.runs if run.text == "Shipped")


def test_names_are_made_safe_and_unknown_files_are_refused(client):
    made = create(filename="../../etc/passwd.md", content="hi")
    assert made["name"] == "passwd.md"
    assert client.get(f"/api/files/{made['id']}/../../secret").status_code == 404
    assert client.get("/api/files/deadbeef/passwd.md").status_code == 404
    with pytest.raises(ValueError, match="kind of file"):
        create(filename="tool.exe", content="MZ")


def test_files_are_always_available(client):
    assert server_states()["files"]["available"] is True


def test_attachments_resolve_to_created_files():
    made = create(filename="notes.txt", content="Remember the milk")
    found = files.attachments(f"{made['id']}")
    assert found[0][0] == "notes.txt" and found[0][1] == b"Remember the milk"
    with pytest.raises(ValueError, match="No file"):
        files.attachments("0123456789abcdef")


def test_old_files_are_cleaned_up_with_the_retention_setting(client):
    import os
    import time

    old, new = create(filename="old.txt", content="x"), create(filename="new.txt", content="y")
    folder = files._own_root() / old["id"]
    stamp = time.time() - 40 * 86400
    os.utime(folder, (stamp, stamp))
    assert files.purge_older_than(30) >= 1
    assert not folder.exists() and (files._own_root() / new["id"]).exists()
    assert files.purge_older_than(0) == 0  # 0 keeps everything


GYM = """| Date | Exercise Name | Sets | Reps | Weight |
|------|---------------|------|------|--------|
| 2026-10-13 | Squat | 3 | 8 | 60.5 |"""


def test_markdown_tables_become_real_spreadsheets(client):
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(download(client, create(filename="gym.xlsx", content=GYM)))).active
    assert [c.value for c in sheet[1]] == ["Date", "Exercise Name", "Sets", "Reps", "Weight"]
    assert [c.value for c in sheet[2]] == ["2026-10-13", "Squat", 3, 8, 60.5]  # numbers stay numbers
    assert sheet.max_row == 2  # no |----| separator row
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2" and sheet.auto_filter.ref
    assert sheet.column_dimensions["B"].width > sheet.column_dimensions["C"].width  # sized to fit
    csv_text = download(client, create(filename="gym.csv", content=GYM)).decode("utf-8-sig")
    assert csv_text.splitlines()[:2] == ["Date,Exercise Name,Sets,Reps,Weight", "2026-10-13,Squat,3,8,60.5"]


def test_markdown_tables_in_word_documents_become_tables(client):
    from docx import Document

    doc = Document(io.BytesIO(download(client, create(filename="plan.docx", content=f"# Plan\n\n{GYM}\n\nKeep going."))))
    assert len(doc.tables) == 1 and [c.text for c in doc.tables[0].rows[1].cells] == ["2026-10-13", "Squat", "3", "8", "60.5"]
    assert not any("|" in p.text for p in doc.paragraphs) and doc.paragraphs[-1].text == "Keep going."


# ---- previews (uploaded and generated files) --------------------------------------------------------

def upload(client, *files):
    response = client.post("/api/files", files=[("file", f) for f in files])
    assert response.status_code == 200, response.text
    return response.json()


def preview(client, made):
    response = client.get(f"{made['url']}/preview")
    assert response.status_code == 200, response.text
    return response.json()


def test_uploaded_files_are_stored_like_generated_ones(client):
    notes, pdf = upload(client, ("../notes.log", b"line one\nline two"), ("plan.pdf", b"%PDF-1.4 fake"))
    assert notes["name"] == "notes.log" and notes["url"] == f"/api/files/{notes['id']}/notes.log"
    assert download(client, notes) == b"line one\nline two"
    assert preview(client, notes) == {"kind": "text", "text": "line one\nline two", "truncated": False}
    assert preview(client, pdf) == {"kind": "pdf"}
    inline = client.get(pdf["url"], params={"inline": 1})
    assert inline.headers["content-disposition"].startswith("inline")
    assert client.post("/api/files", files=[("file", ("big.txt", b"x" * (files.MAX_UPLOAD + 1)))]).status_code == 422


def test_previews_show_tables_documents_and_text(client):
    sheet = preview(client, create(filename="gym.xlsx", content=GYM))
    assert sheet["kind"] == "table" and sheet["sheets"][0]["rows"] == [
        ["Date", "Exercise Name", "Sets", "Reps", "Weight"], ["2026-10-13", "Squat", 3, 8, 60.5]]
    assert preview(client, create(filename="gym.csv", content=GYM))["sheets"][0]["rows"][1][1] == "Squat"
    word = preview(client, create(filename="plan.docx", content=f"# Plan\n\n{GYM}\n\nKeep going."))
    assert word["kind"] == "markdown" and word["text"].startswith("# Plan") and "| 2026-10-13 | Squat |" in word["text"]
    assert word["text"].rstrip().endswith("Keep going.")
    assert preview(client, create(filename="notes.md", content="# Hi"))["kind"] == "markdown"
    assert preview(client, create(filename="data.json", content='{"a":1}'))["kind"] == "text"
    assert preview(client, create(filename="page.html", content="<b>hi</b>")) == {"kind": "html", "text": "<b>hi</b>", "truncated": False}
    page = create(filename="page.html", content="<script>alert(1)</script>")
    assert client.get(page["url"], params={"inline": 1}).headers["content-disposition"].startswith("attachment")  # never run it here


def test_long_tables_and_unknown_files_are_previewed_safely(client):
    big = preview(client, create(filename="big.csv", content="n\n" + "\n".join(map(str, range(500)))))
    assert len(big["sheets"][0]["rows"]) == files.PREVIEW_ROWS and big["sheets"][0]["truncated"] is True
    (blob,) = upload(client, ("photo.bin", bytes(range(256))))
    assert preview(client, blob) == {"kind": "none"}
