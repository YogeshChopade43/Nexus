"""Shared test data builders."""
from __future__ import annotations

import io


def make_pdf(pages: list[str]) -> bytes:
    """A minimal PDF with one page per string; an empty string makes a page with no text layer."""
    objects: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>", b""]  # pages object filled in below
    kids = []
    for text in pages:
        page_no = len(objects) + 1
        content_no = page_no + 1
        kids.append(f"{page_no} 0 R")
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 %d 0 R >> >> >>" % (content_no, 3 + 2 * len(pages))
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offsets)
    return out + b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)


def make_docx(heading: str, paragraph: str, table: list[list[str]]) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(paragraph)
    grid = doc.add_table(rows=len(table), cols=len(table[0]))
    for r, row in enumerate(table):
        for c, cell in enumerate(row):
            grid.cell(r, c).text = cell
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
